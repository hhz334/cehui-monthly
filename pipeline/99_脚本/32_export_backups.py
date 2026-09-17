# -*- coding: utf-8 -*-
"""导出对照件：全量底稿（68 条）、领导终稿（23 条），并写入 xlsx 工作表。"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from period_config import P  # noqa: E402
from fetch_utils import RAW, ROOT, load_json  # noqa: E402

BOARDS = ["政策法规", "技术应用类", "科技前沿类"]      # 2026-09-17 起取消“媒体动态类”
FULL = os.path.join(RAW, "full68_selection.json")
LEADER = os.path.join(RAW, "leader_%s.json" % (P.leader_label or "0917"))
XLSX = os.path.join(ROOT, "00_资讯目录.xlsx")


def main():
    full = load_json(FULL)
    leader = load_json(LEADER)

    lines = [f"# 全量底稿（{P.title_range}）", "",
             f"> 说明：本文件是本期的全量候选（目录与领导终稿均由此产生），供追溯与下期备选。",
             f"- 条数：{len(full)}", ""]
    for b in BOARDS:
        sub = [s for s in full if s.get("board") == b]
        if not sub:
            continue
        lines += [f"## {b}（{len(sub)} 条）", ""]
        for i, s in enumerate(sub, 1):
            lines.append(f"{i}. {s.get('title','')}")
            lines.append(f"   - 来源：{s.get('unit','')}｜{s.get('date_override','')}｜业务条线：{s.get('business','')}")
            lines.append(f"   - 链接：{s.get('source_url_checked','')}")
        lines.append("")
    open(os.path.join(ROOT, "_备查", f"全量底稿（{P.title_range}）.md"), "w", encoding="utf-8").write("\n".join(lines) + "\n")

    lines = [f"# 领导终稿（{P.leader_label or '0917'}，{len(leader['final'])} 条）", "",
             "> 说明：领导批注后的定稿，作为下一轮选目的骨架。",
             f"- 条数：{len(leader['final'])}｜删除：{leader.get('removed_count', 0)}｜新增：{len(leader.get('added', []))}", ""]
    n = 0
    for b in BOARDS:
        sub = [x for x in leader["final"] if x["sec"] == b]
        if not sub:
            continue
        lines += [f"## {b}（{len(sub)} 条）", ""]
        for x in sub:
            n += 1
            lines.append(f"{n}. {x['title']}")
            if x.get("src"):
                lines.append(f"   - {x['src']}")
            for u in (x.get("urls") or [])[:2]:
                lines.append(f"   - {u}")
        lines.append("")
    leader_md = os.path.join(ROOT, "_备查", f"领导终稿（{len(leader['final'])}条）.md")
    open(leader_md, "w", encoding="utf-8").write("\n".join(lines) + "\n")

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = load_workbook(XLSX)
    name = "全量底稿"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    cols = ["序号", "板块", "标题", "发布单位", "属地层级", "省份", "业务条线", "优先级", "发布日期", "来源链接"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5597")
    for i, s in enumerate(full, 1):
        ws.append([i, s.get("board", ""), s.get("title", ""), s.get("unit", ""), s.get("level", ""),
                   s.get("province", ""), s.get("business", ""), s.get("priority", ""),
                   s.get("date_override", ""), s.get("source_url_checked", "")])
    for col, w in {"A": 6, "B": 12, "C": 56, "D": 24, "E": 10, "F": 8, "G": 22, "H": 8, "I": 12, "J": 60}.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column_letter in "CDJ")
    ws.freeze_panes = "A2"
    wb.save(XLSX)
    print(f"全量底稿 {len(full)} 条 → _备查/全量底稿（{P.title_range}）.md + xlsx 工作表「全量底稿」")
    print(f"领导终稿 {len(leader['final'])} 条 → {os.path.basename(leader_md)}")


if __name__ == "__main__":
    main()
