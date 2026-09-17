# -*- coding: utf-8 -*-
"""生成《测绘地理信息周讯》线索池：82—87期全部条目及其回溯结果。

新口径下，周讯只作选题线索，条目须回溯到原发布方官网或白名单公众号原文才能进目录。
输出：05_周讯线索池.md，并向 00_资讯目录.xlsx 追加“周讯线索池”工作表。
"""
import json
import os
import re
import sys
from difflib import SequenceMatcher

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from fetch_utils import ROOT, load_json  # noqa: E402
from period_config import P  # noqa: E402

WEEKLY_JSON = os.path.join(ROOT, "03_参考_山东周讯", "周讯条目提取.json")
SEL_BOARDS = os.path.join(P.sel_dir, "catalog_selection_boards.json")
SEL = SEL_BOARDS if os.path.exists(SEL_BOARDS) else os.path.join(P.sel_dir, "catalog_selection.json")
XLSX = os.path.join(ROOT, "00_资讯目录.xlsx")
OUT_MD = os.path.join(ROOT, "05_周讯线索池.md")


def norm(text):
    return re.sub(r"[\s\W_]+", "", text or "").lower()


def similar(a, b):
    """整串相似度（SequenceMatcher 比值），比集合覆盖率更能避免“同尾词误配”。"""
    a, b = norm(a), norm(b)
    if not a or not b:
        return 0.0
    return SequenceMatcher(None, a, b).ratio()


def quoted(text):
    """取出标题中被《》引用的文件名，用于同名文件的跨来源匹配。"""
    m = re.search(r"《([^》]+)》", text or "")
    return norm(m.group(1)) if m else ""


def is_same(a, b, ratio):
    """判定周讯条目与目录条目是否同一事件。"""
    if ratio >= 0.80:
        return True
    qa, qb = quoted(a), quoted(b)
    return bool(qa) and qa == qb and ratio >= 0.55


def main():
    weekly = load_json(WEEKLY_JSON)
    selection = json.load(open(SEL, encoding="utf-8"))
    titles = [s["title"] for s in selection]

    rows = []
    for w in weekly:
        best_score, best_title = 0.0, ""
        for t in titles:
            score = similar(w["title"], t)
            if score > best_score or is_same(w["title"], t, score):
                best_score, best_title = score, t
        matched = bool(best_title) and is_same(w["title"], best_title, best_score)
        rows.append({**w, "入选": matched, "对应目录条目": best_title if matched else "",
                     "相似度": round(best_score, 2)})

    lines = ["# 《测绘地理信息周讯》82—87 期线索池", "",
             "> 口径：周讯只作选题线索，不作为来源。条目须回溯到原发布方官网或白名单公众号原文后才能进目录；",
             "> 本表列出 82—87 期全部条目（由整期 PDF 文本层提取）与其回溯结果，"
             "已回溯的原发布方条目见《资讯目录》，未回溯的留在本表继续跟踪。", ""]
    for issue in [82, 83, 84, 85, 86, 87]:
        part = [r for r in rows if r["issue"] == issue]
        picked = sum(1 for r in part if r["入选"])
        lines += [f"## 第{issue}期（共 {len(part)} 条，已回溯 {picked} 条）", "",
                  "| # | 栏目 | 标题 | 原始来源 | 回溯结果 | 对应目录条目 |",
                  "| --- | --- | --- | --- | --- | --- |"]
        for r in part:
            lines.append(f"| {r['no']} | {r['section']} | {r['title']} | {r['source'] or '未标注'} | "
                         f"{'已回溯入目录' if r['入选'] else '待回溯'} | {r['对应目录条目']} |")
        lines.append("")
    with open(OUT_MD, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))
    print(f"已生成 {OUT_MD}（{len(rows)} 条，已回溯 {sum(1 for r in rows if r['入选'])} 条）")

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = load_workbook(XLSX)
    for legacy in ("周讯延伸池", "周讯线索池"):
        if legacy in wb.sheetnames:
            del wb[legacy]
    ws = wb.create_sheet("周讯线索池")
    cols = ["期号", "周讯栏目", "标题", "原始来源", "回溯结果", "对应目录条目", "相似度"]
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DCE6F1")
    for r in sorted(rows, key=lambda x: (x["issue"], x["no"])):
        ws.append([f"第{r['issue']}期", r["section"], r["title"], r["source"] or "未标注",
                   "已回溯入目录" if r["入选"] else "待回溯", r["对应目录条目"], r["相似度"]])
    for col, width in {"A": 10, "B": 18, "C": 56, "D": 26, "E": 16, "F": 40, "G": 8}.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    wb.save(XLSX)
    print(f"已向 {os.path.basename(XLSX)} 追加工作表“周讯线索池”")


if __name__ == "__main__":
    main()
